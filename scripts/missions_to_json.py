"""把欧委会公布的 Séjourné 出差报销表(xlsx)转成 data/sejourne_missions.json,供「✈️ 出差地图」页读取。

用法:
    python3 scripts/missions_to_json.py Missions_of_Executive_Vice-President_Stephane_Sejourne_published_on_XX_XX_XXXX.xlsx

报销表在他委员主页 Transparency → Missions 下载,每两个月更新一次,每份都是从上任起的全量。
幂等:按出发日期合并。已有条目里手写的字段原样保留——中文摘要(paris/other)、日历确认的巴黎日期(calp)、
修正过的路线(stops_fixed=true 时的 stops);只刷新费用、事由原文等官方字段。
新出现的条目先把法语事由放进 other 顶着(fr=true),中文摘要要手动补进 JSON。
Location 里出现 utils/missions.py 的 CITIES 没有的地名会直接报错,先补坐标再跑。
"""
import json
import pathlib
import re
import sys
from datetime import datetime

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from utils.missions import CITIES, DATA_PATH  # noqa: E402

# 手写字段:重跑时不覆盖
MANUAL_FIELDS = ("paris", "other", "fr", "calp", "note")


def _num(v):
    return round(float(v or 0), 2)


def read_xlsx(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    title = ws.cell(1, 1).value or ""
    rows = []
    for r in ws.iter_rows(values_only=True):
        if not (r[0] and re.match(r"\d\d/\d\d/\d{4}$", str(r[0]).strip())):
            continue
        rows.append({
            "f": datetime.strptime(r[0].strip(), "%d/%m/%Y").strftime("%Y-%m-%d"),
            "t": datetime.strptime(str(r[1]).strip(), "%d/%m/%Y").strftime("%Y-%m-%d"),
            "stops": [s.strip() for s in str(r[2]).split(",") if s.strip()],
            "purpose_fr": re.sub(r"\s+", " ", str(r[3] or "")).strip(),
            "travel": _num(r[5]), "hotel": _num(r[6]), "daily": _num(r[7]), "misc": _num(r[8]),
            "taxi": "air taxi" in str(r[9] or "").lower(),
        })
    return title, rows


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    title, rows = read_xlsx(sys.argv[1])
    if not rows:
        sys.exit("❌ 表里没读到任何出差行(第一列应为 DD/MM/YYYY 日期),检查文件是否是 Missions 报销表")

    missing = sorted({s for r in rows for s in r["stops"]} - CITIES.keys())
    if missing:
        sys.exit(f"❌ 这些地名没有坐标,先在 utils/missions.py 的 CITIES 里补上再跑:{missing}")

    old = {}
    if DATA_PATH.exists():
        old = {m["f"]: m for m in json.loads(DATA_PATH.read_text(encoding="utf-8"))["missions"]}

    merged, added = [], 0
    for r in rows:
        prev = old.get(r["f"])
        if prev:
            for k in MANUAL_FIELDS:
                if k in prev:
                    r[k] = prev[k]
            if prev.get("stops_fixed"):
                r["stops"], r["stops_fixed"] = prev["stops"], True
        else:
            added += 1
            r.update(paris=None, other=r["purpose_fr"], fr=True, calp=[])
        merged.append(r)
    merged.sort(key=lambda m: m["f"])

    DATA_PATH.parent.mkdir(exist_ok=True)
    DATA_PATH.write_text(json.dumps({"source": title.strip(), "missions": merged},
                                    ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"✅ {len(merged)} 条写入 {DATA_PATH.relative_to(DATA_PATH.parent.parent)}(新增 {added} 条,"
          f"范围 {merged[0]['f']} ~ {merged[-1]['f']})")
    if added:
        print("   新增条目的 other 暂为法语原文(fr=true),记得补中文摘要")


if __name__ == "__main__":
    main()
